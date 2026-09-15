"""Excel report generation for CFD simulations."""

import json
import tempfile
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter

from core.logging import get_logger

logger = get_logger(__name__)


# ============================================================
# STYLING CONSTANTS
# ============================================================

PRIMARY_COLOR = "3B82F6"
HEADER_BG = "1E40AF"
HEADER_FG = "FFFFFF"
LABEL_BG = "DBEAFE"
LABEL_FG = "1E3A8A"
SUCCESS_COLOR = "10B981"
WARNING_COLOR = "F59E0B"
ERROR_COLOR = "EF4444"

# Fonts
TITLE_FONT = Font(name="Calibri", size=18, bold=True, color="1F2937")
SUBTITLE_FONT = Font(name="Calibri", size=12, italic=True, color="6B7280")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color=HEADER_FG)
LABEL_FONT = Font(name="Calibri", size=11, bold=True, color=LABEL_FG)
VALUE_FONT = Font(name="Calibri", size=11, color="1F2937")

# Borders
THIN_BORDER = Border(
    left=Side(style="thin", color="E5E7EB"),
    right=Side(style="thin", color="E5E7EB"),
    top=Side(style="thin", color="E5E7EB"),
    bottom=Side(style="thin", color="E5E7EB"),
)

# Fills
HEADER_FILL = PatternFill("solid", fgColor=HEADER_BG)
LABEL_FILL = PatternFill("solid", fgColor=LABEL_BG)
ALT_ROW_FILL = PatternFill("solid", fgColor="F9FAFB")


# ============================================================
# HELPERS
# ============================================================

def _write_title(ws, row: int, title: str, subtitle: str = "") -> int:
    """Write a title with optional subtitle. Returns next row."""
    ws.cell(row=row, column=1, value=title).font = TITLE_FONT
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    row += 1

    if subtitle:
        ws.cell(row=row, column=1, value=subtitle).font = SUBTITLE_FONT
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        row += 1

    row += 1  # Spacer
    return row


def _write_key_value_section(
    ws,
    start_row: int,
    title: str,
    data: Dict[str, Any],
) -> int:
    """Write a section with key-value pairs."""
    row = start_row

    # Section header
    ws.cell(row=row, column=1, value=title).font = HEADER_FONT
    ws.cell(row=row, column=1).fill = HEADER_FILL
    ws.cell(row=row, column=2).fill = HEADER_FILL
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    row += 1

    # Key-value rows
    for i, (key, value) in enumerate(data.items()):
        label_cell = ws.cell(row=row, column=1, value=str(key))
        label_cell.font = LABEL_FONT
        label_cell.fill = LABEL_FILL
        label_cell.border = THIN_BORDER

        # Format value
        if value is None:
            display = "N/A"
        elif isinstance(value, float):
            if abs(value) < 0.001 or abs(value) > 1e6:
                display = f"{value:.4e}"
            else:
                display = f"{value:,.4f}"
        elif isinstance(value, int):
            display = f"{value:,}"
        elif isinstance(value, (list, dict)):
            display = json.dumps(value)
        else:
            display = str(value)

        value_cell = ws.cell(row=row, column=2, value=display)
        value_cell.font = VALUE_FONT
        value_cell.border = THIN_BORDER

        if i % 2 == 1:
            value_cell.fill = ALT_ROW_FILL

        row += 1

    return row + 1


def _write_table(
    ws,
    start_row: int,
    title: str,
    headers: List[str],
    rows: List[List[Any]],
) -> int:
    """Write a table with headers and rows."""
    row = start_row

    # Title
    if title:
        ws.cell(row=row, column=1, value=title).font = Font(bold=True, size=12)
        ws.merge_cells(
            start_row=row, start_column=1,
            end_row=row, end_column=len(headers),
        )
        row += 1

    # Headers
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")
    row += 1

    # Rows
    for i, data_row in enumerate(rows):
        for col, value in enumerate(data_row, start=1):
            if value is None:
                display = "N/A"
            elif isinstance(value, float):
                if abs(value) < 0.001 or abs(value) > 1e6:
                    display = f"{value:.4e}"
                else:
                    display = f"{value:,.4f}"
            elif isinstance(value, int):
                display = f"{value:,}"
            else:
                display = str(value)

            cell = ws.cell(row=row, column=col, value=display)
            cell.font = VALUE_FONT
            cell.border = THIN_BORDER
            if i % 2 == 1:
                cell.fill = ALT_ROW_FILL
        row += 1

    return row + 1


def _auto_column_widths(ws, min_width: int = 15, max_width: int = 60) -> None:
    """Auto-adjust column widths based on content."""
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        width = min(max(max_len + 2, min_width), max_width)
        ws.column_dimensions[col_letter].width = width


# ============================================================
# MAIN REPORT GENERATOR
# ============================================================

class ExcelReportGenerator:
    """Generate a professional Excel report from a CFD simulation."""

    def __init__(
        self,
        simulation: Dict[str, Any],
        project: Dict[str, Any],
        geometry: Optional[Dict[str, Any]] = None,
        logs: Optional[List[Dict[str, Any]]] = None,
    ):
        self.simulation = simulation
        self.project = project
        self.geometry = geometry or {}
        self.logs = logs or []

        # Parse JSON config
        try:
            self.config = json.loads(simulation.get("config") or "{}")
        except Exception:
            self.config = {}

        try:
            self.metrics = json.loads(simulation.get("metrics") or "{}")
        except Exception:
            self.metrics = {}

        try:
            self.available_fields = json.loads(
                simulation.get("available_fields") or "[]"
            )
        except Exception:
            self.available_fields = []

    def generate(self, output_path: Optional[Path] = None) -> Path:
        """Generate the complete Excel report."""
        wb = Workbook()
        wb.remove(wb.active)  # Remove default sheet

        # Build all sheets
        self._create_summary_sheet(wb)
        self._create_geometry_sheet(wb)
        self._create_fluid_sheet(wb)
        self._create_boundary_sheet(wb)
        self._create_solver_sheet(wb)
        self._create_mesh_sheet(wb)
        self._create_results_sheet(wb)
        self._create_ai_sheet(wb)
        self._create_logs_sheet(wb)

        # Save
        if output_path is None:
            with tempfile.NamedTemporaryFile(
                suffix=".xlsx", delete=False
            ) as tmp:
                output_path = Path(tmp.name)

        wb.save(str(output_path))
        logger.info(f"Excel report saved: {output_path}")
        return output_path

    # --------------------------------------------------------
    # SHEET 1: SUMMARY
    # --------------------------------------------------------

    def _create_summary_sheet(self, wb: Workbook) -> None:
        ws = wb.create_sheet("Summary")
        row = 1

        row = _write_title(
            ws, row,
            "CFD Simulation Report",
            f"Generated on {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        )

        # Project info
        row = _write_key_value_section(ws, row, "Project Information", {
            "Project Name": self.project.get("name"),
            "Project ID": self.project.get("project_id"),
            "CFD Type": self.project.get("cfd_type"),
            "Description": self.project.get("description") or "N/A",
            "Created": str(self.project.get("created_at", ""))[:19],
        })

        # Simulation info
        row = _write_key_value_section(ws, row, "Simulation Information", {
            "Simulation ID": self.simulation.get("simulation_id"),
            "Status": self.simulation.get("status"),
            "Solver": self.simulation.get("solver"),
            "Started": str(self.simulation.get("started_at", ""))[:19] or "N/A",
            "Completed": str(self.simulation.get("completed_at", ""))[:19] or "N/A",
            "Execution Time": f"{self.simulation.get('execution_time', 0):.2f} s",
            "Demo Mode": "Yes" if self.simulation.get("is_demo") else "No",
        })

        # Key results
        row = _write_key_value_section(ws, row, "Key Results", {
            "Max Velocity (m/s)": self.metrics.get("max_velocity"),
            "Avg Velocity (m/s)": self.metrics.get("avg_velocity"),
            "Max Pressure (Pa)": self.metrics.get("max_pressure"),
            "Min Pressure (Pa)": self.metrics.get("min_pressure"),
            "Pressure Drop (Pa)": self.metrics.get("pressure_drop"),
            "Mesh Cells": self.metrics.get("mesh_cells"),
            "Converged": self.metrics.get("converged"),
            "Iterations": self.metrics.get("iterations"),
        })

        _auto_column_widths(ws)

    # --------------------------------------------------------
    # SHEET 2: GEOMETRY
    # --------------------------------------------------------

    def _create_geometry_sheet(self, wb: Workbook) -> None:
        ws = wb.create_sheet("Geometry")
        row = 1

        row = _write_title(ws, row, "Geometry Information")

        if not self.geometry:
            ws.cell(row=row, column=1, value="No geometry information available.")
            return

        row = _write_key_value_section(ws, row, "File", {
            "Filename": self.geometry.get("filename"),
            "Original Filename": self.geometry.get("original_filename"),
            "Format": self.geometry.get("format"),
            "File Size (bytes)": self.geometry.get("file_size"),
            "Validation Status": self.geometry.get("validation_status"),
            "Validation Message": self.geometry.get("validation_message") or "OK",
            "Uploaded": str(self.geometry.get("created_at", ""))[:19],
        })

        _auto_column_widths(ws)

    # --------------------------------------------------------
    # SHEET 3: FLUID PROPERTIES
    # --------------------------------------------------------

    def _create_fluid_sheet(self, wb: Workbook) -> None:
        ws = wb.create_sheet("Fluid Properties")
        row = 1

        row = _write_title(ws, row, "Fluid Properties")

        fluid = self.config.get("fluid", {})
        density = fluid.get("density", 0)
        viscosity = fluid.get("dynamic_viscosity", 0)
        kinematic = viscosity / density if density > 0 else 0

        row = _write_key_value_section(ws, row, "Properties", {
            "Fluid": fluid.get("name", "N/A"),
            "Density (kg/m³)": density,
            "Dynamic Viscosity (Pa·s)": viscosity,
            "Kinematic Viscosity (m²/s)": kinematic,
            "Temperature (K)": fluid.get("temperature"),
            "Reference Pressure (Pa)": fluid.get("pressure"),
        })

        _auto_column_widths(ws)

    # --------------------------------------------------------
    # SHEET 4: BOUNDARY CONDITIONS
    # --------------------------------------------------------

    def _create_boundary_sheet(self, wb: Workbook) -> None:
        ws = wb.create_sheet("Boundary Conditions")
        row = 1

        row = _write_title(ws, row, "Boundary Conditions")

        bc = self.config.get("boundary_conditions", {})
        inlet = bc.get("inlet", {})
        outlet = bc.get("outlet", {})
        walls = bc.get("walls", {})

        headers = ["Boundary", "Field", "Value", "Unit"]
        rows = [
            ["Inlet", "Velocity", inlet.get("velocity"), "m/s"],
            ["Inlet", "Pressure", inlet.get("pressure"), "Pa"],
            ["Inlet", "Temperature", inlet.get("temperature"), "K"],
            ["Outlet", "Pressure", outlet.get("pressure"), "Pa"],
            ["Walls", "Type", walls.get("type"), "-"],
        ]

        row = _write_table(ws, row, "Specified Conditions", headers, rows)

        _auto_column_widths(ws)

    # --------------------------------------------------------
    # SHEET 5: SOLVER
    # --------------------------------------------------------

    def _create_solver_sheet(self, wb: Workbook) -> None:
        ws = wb.create_sheet("Solver")
        row = 1

        row = _write_title(ws, row, "Solver Configuration")

        solver = self.config.get("solver_settings", {})

        row = _write_key_value_section(ws, row, "Settings", {
            "Solver": solver.get("solver", "N/A"),
            "Max Iterations": solver.get("max_iterations"),
            "Convergence Tolerance": solver.get("convergence_tolerance"),
            "Time Step (s)": solver.get("time_step"),
        })

        _auto_column_widths(ws)

    # --------------------------------------------------------
    # SHEET 6: MESH
    # --------------------------------------------------------

    def _create_mesh_sheet(self, wb: Workbook) -> None:
        ws = wb.create_sheet("Mesh")
        row = 1

        row = _write_title(ws, row, "Mesh Configuration")

        mesh = self.config.get("mesh_settings", {})

        row = _write_key_value_section(ws, row, "Parameters", {
            "Base Cell Size (m)": mesh.get("base_cell_size"),
            "Refinement Level": mesh.get("refinement_level"),
            "Boundary Layers": mesh.get("boundary_layers"),
            "Estimated Cells": mesh.get("estimated_cells"),
            "Actual Cells": self.metrics.get("mesh_cells"),
        })

        _auto_column_widths(ws)

    # --------------------------------------------------------
    # SHEET 7: RESULTS
    # --------------------------------------------------------

    def _create_results_sheet(self, wb: Workbook) -> None:
        ws = wb.create_sheet("Results")
        row = 1

        row = _write_title(ws, row, "Simulation Results")

        headers = ["Metric", "Value", "Unit"]
        rows = [
            ["Maximum Velocity", self.metrics.get("max_velocity"), "m/s"],
            ["Minimum Velocity", self.metrics.get("min_velocity"), "m/s"],
            ["Average Velocity", self.metrics.get("avg_velocity"), "m/s"],
            ["Maximum Pressure", self.metrics.get("max_pressure"), "Pa"],
            ["Minimum Pressure", self.metrics.get("min_pressure"), "Pa"],
            ["Pressure Drop", self.metrics.get("pressure_drop"), "Pa"],
            ["Mesh Cells", self.metrics.get("mesh_cells"), "cells"],
            ["Execution Time", self.simulation.get("execution_time"), "s"],
            ["Iterations", self.metrics.get("iterations"), "-"],
            ["Converged", "Yes" if self.metrics.get("converged") else "No", "-"],
        ]

        row = _write_table(ws, row, "Computed Metrics", headers, rows)

        # Available fields
        if self.available_fields:
            row += 1
            row = _write_table(
                ws, row, "Available Fields",
                ["Field"],
                [[f] for f in self.available_fields],
            )

        _auto_column_widths(ws)

    # --------------------------------------------------------
    # SHEET 8: AI ANALYSIS
    # --------------------------------------------------------

    def _create_ai_sheet(self, wb: Workbook) -> None:
        ws = wb.create_sheet("AI Analysis")
        row = 1

        row = _write_title(ws, row, "AI Analysis")

        ws.cell(row=row, column=1, value="Status").font = LABEL_FONT
        ws.cell(row=row, column=2, value="No trained model configured").font = VALUE_FONT
        row += 1

        ws.cell(row=row, column=1, value="Model").font = LABEL_FONT
        ws.cell(row=row, column=2, value="N/A").font = VALUE_FONT
        row += 1

        ws.cell(row=row, column=1, value="Note").font = LABEL_FONT
        ws.cell(
            row=row, column=2,
            value="AI predictions will appear here once a model is trained.",
        ).font = VALUE_FONT

        _auto_column_widths(ws)

    # --------------------------------------------------------
    # SHEET 9: LOGS
    # --------------------------------------------------------

    def _create_logs_sheet(self, wb: Workbook) -> None:
        ws = wb.create_sheet("Simulation Logs")
        row = 1

        row = _write_title(ws, row, "Simulation Logs")

        if not self.logs:
            ws.cell(row=row, column=1, value="No logs available.")
            return

        headers = ["Timestamp", "Type", "Content"]
        rows = [
            [
                str(log.get("timestamp", ""))[:19],
                log.get("log_type", "info"),
                log.get("content", ""),
            ]
            for log in self.logs
        ]

        _write_table(ws, row, "", headers, rows)
        _auto_column_widths(ws, min_width=20, max_width=120)


# ============================================================
# CONVENIENCE FUNCTION
# ============================================================

def generate_report_for_simulation(simulation_id: str) -> Optional[Path]:
    """Generate an Excel report for a simulation by ID."""
    from storage import get_geometry, get_logs, get_project, get_simulation

    sim = get_simulation(simulation_id)
    if not sim:
        logger.error(f"Simulation not found: {simulation_id}")
        return None

    project = get_project(sim["project_id"])
    if not project:
        logger.error(f"Project not found: {sim['project_id']}")
        return None

    geometry = get_geometry(sim["project_id"])
    logs = get_logs(simulation_id)

    generator = ExcelReportGenerator(
        simulation=sim,
        project=project,
        geometry=geometry,
        logs=logs,
    )

    return generator.generate()